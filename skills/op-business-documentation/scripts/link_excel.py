# link_excel.py <PREFIX> <workbook.xlsx> [<workbook.xlsx> ...]   (needs openpyxl)
# Pass every Registries workbook of the product (Shared and all features) in one run.
import os, re, sys
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
prefix, paths = sys.argv[1], [os.path.abspath(p) for p in sys.argv[2:]]
tok = re.compile(re.escape(prefix) + r'-[A-Z0-9]+(?:-[A-Z0-9]+)+')
books = {p: load_workbook(p) for p in paths}
home = {}  # ID -> (workbook path, sheet, row, last column)
for p, wb in books.items():
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible': continue
        last = get_column_letter(ws.max_column)
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if isinstance(v, str) and tok.fullmatch(v.strip()):
                i = v.strip()
                if i in home: print('DUPLICATE ID', i, os.path.basename(home[i][0]), os.path.basename(p))
                home[i] = (p, ws.title, r, last)
for p, wb in books.items():
    for k in [n for n in wb.defined_names if n.startswith(prefix + '_')]:
        del wb.defined_names[k]
    for i, (hp, t, r, last) in home.items():
        if hp != p: continue
        n = i.replace('-', '_')
        wb.defined_names[n] = DefinedName(n, attr_text=f"'{t}'!$A${r}:${last}${r}")
broken, multi = [], 0
for p, wb in books.items():
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if not isinstance(c.value, str): continue
                ids = tok.findall(c.value)
                if c.column == 1 and ids and tok.fullmatch(c.value.strip()): continue
                broken += [(os.path.basename(p), ws.title, c.coordinate, i) for i in ids if i not in home]
                good = [i for i in ids if i in home]
                if not good: c.hyperlink = None; continue
                first = good[0]; name = first.replace('-', '_'); hp = home[first][0]
                if hp == p:
                    c.hyperlink = Hyperlink(ref=c.coordinate, location=name, display=c.value)
                else:
                    c.hyperlink = os.path.relpath(hp, os.path.dirname(p)).replace(os.sep, '/') + '#' + name
                c.font = Font(name=c.font.name, size=c.font.size, color='1155CC', underline='single')
                multi += len(good) > 1
    wb.save(p)
print(len(home), 'IDs defined across', len(paths), 'workbooks;', multi, 'cells with several IDs')
for b in broken: print('BROKEN REFERENCE', *b)
